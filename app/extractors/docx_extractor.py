"""
MedExtract-IA: DOCX Layout Extractor

Extracts structured markdown text + images from .docx files using python-docx.
Preserves:
  - Bold text as **markdown** markers
  - Tables as markdown/HTML
  - Images as in-memory BytesIO buffers
"""

import io
import re
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def _extract_rich_text(element) -> str:
    """Extract text from a paragraph or cell while preserving inline bold and italic formatting."""
    parts = []
    
    # Check if the element has 'runs' (like a Paragraph)
    if hasattr(element, 'runs'):
        for run in element.runs:
            text = run.text
            if not text:
                continue
            
            # Preserve leading/trailing spaces outside the markdown tags
            lspace = " " * (len(text) - len(text.lstrip(" ")))
            rspace = " " * (len(text) - len(text.rstrip(" ")))
            stripped = text.strip(" ")
            
            if stripped:
                if run.bold:
                    stripped = f"**{stripped}**"
                elif run.italic:
                    stripped = f"*{stripped}*"
            
            parts.append(lspace + stripped + rspace)
            
    # Check if the element has 'paragraphs' (like a Cell)
    elif hasattr(element, 'paragraphs'):
        para_texts = []
        for p in element.paragraphs:
            p_text = _extract_rich_text(p)
            if p_text:
                para_texts.append(p_text)
        return "<br>".join(para_texts)
        
    line = "".join(parts)
    # Clean up adjacent bold markers: **text****more** → **text more**
    line = re.sub(r'\*\*\s*\*\*', ' ', line)
    return line


def extract_docx(file_bytes: bytes) -> tuple[str, dict[str, io.BytesIO]]:
    """
    Extract text and images from a .docx file.

    Args:
        file_bytes: Raw bytes of the .docx file

    Returns:
        Tuple of (markdown_text, image_buffers)
        - markdown_text: Clean markdown with **bold** preserved and [[IMG_N]] placeholders
        - image_buffers: Dict mapping image key (e.g. 'img_1') to BytesIO buffer
    """
    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    image_buffers: dict[str, io.BytesIO] = {}
    image_counter = 0

    # ─── Build image map from document relationships ─────────────────
    image_rels: dict[str, bytes] = {}
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                image_rels[rel.rId] = rel.target_part.blob
            except Exception:
                pass

    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    # ─── Process elements sequentially ───────────────────────────────
    for block in doc.element.body:
        if isinstance(block, CT_P):
            para = Paragraph(block, doc)
            # Check for inline images in the paragraph's XML
            para_xml = para._element.xml
            if "blip" in para_xml or "imagedata" in para_xml:
                # Extract image references
                for rel_id in re.findall(r'r:embed="([^"]+)"', para_xml):
                    if rel_id in image_rels:
                        image_counter += 1
                        key = f"img_{image_counter}"
                        buffer = io.BytesIO(image_rels[rel_id])
                        image_buffers[key] = buffer

                if not para.text.strip():
                    lines.append(f"[[IMG_{image_counter}]]")
                    continue

            # Extract text with run-level bold detection
            line = _extract_rich_text(para).strip()
            if line:
                lines.append(line)

        elif isinstance(block, CT_Tbl):
            table = Table(block, doc)
            table_lines = []
            for i, row in enumerate(table.rows):
                cells = []
                for cell in row.cells:
                    # Extract rich text from the cell (preserves bold/italics and handles paragraphs with <br>)
                    cell_text = _extract_rich_text(cell).strip()
                    
                    # Check for images in table cell
                    cell_xml = cell._element.xml
                    if "blip" in cell_xml or "imagedata" in cell_xml:
                        for rel_id in re.findall(r'r:embed="([^"]+)"', cell_xml):
                            if rel_id in image_rels:
                                image_counter += 1
                                key = f"img_{image_counter}"
                                buffer = io.BytesIO(image_rels[rel_id])
                                image_buffers[key] = buffer
                                cell_text += f" [[IMG_{image_counter}]]"
                    
                    # Escape pipe characters in cell text to avoid breaking markdown tables
                    cell_text = cell_text.replace("|", "\\|")
                    cells.append(cell_text)

                table_lines.append("| " + " | ".join(cells) + " |")
                if i == 0:
                    table_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
            lines.append("\n".join(table_lines))
            lines.append("")  # blank line after table

    markdown = "\n".join(lines)

    # ─── Normalize whitespace while preserving structure ─────────────
    markdown = re.sub(r'\n{3,}', '\n\n', markdown)

    return markdown, image_buffers
